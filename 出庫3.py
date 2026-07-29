import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side
)
from openpyxl.utils import get_column_letter


in_file_path = r"\\dionas\管理部\備品管理\入出庫履歴\2026入庫履歴.xlsx"
out_file_path = r"\\dionas\管理部\備品管理\入出庫履歴\2026出庫履歴.xlsx"
code_sheet_name = "コード別集計"
department_sheet_name = "部門別集計"


def normalize_code(value):
    if pd.isna(value):
        return ""
    code = str(value).strip()
    if code.endswith(".0"):
        code = code[:-2]
    return code


header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
header_font = Font(bold=True)
total_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
total_font = Font(bold=True)
thin_side = Side(style="thin", color="808080")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)


# ======== 入庫記録 ========
in_df = pd.read_excel(in_file_path)
in_df["コード"] = (in_df["コード"].apply(normalize_code))
in_df["数量"] = pd.to_numeric(in_df["数量"],errors="coerce").fillna(0)
in_df["金額"] = pd.to_numeric(in_df["金額"],errors="coerce").fillna(0)

price_summary = (
    in_df.groupby("コード", as_index=False, dropna=False)
    .agg(入庫総数量=("数量", "sum"), 入庫総金額=("金額", "sum"))
)

price_summary["総単価"] = np.where(
    price_summary["入庫総数量"] != 0,
    (price_summary["入庫総金額"] / price_summary["入庫総数量"]),
    np.nan
)

price_summary["総単価"] = price_summary["総単価"].round(1)
price_map = price_summary[["コード", "総単価"]].copy()


# ======== 出庫記録 ========
out_df = pd.read_excel(out_file_path)
out_df["コード"] = (out_df["コード"].apply(normalize_code))
out_df["数量"] = pd.to_numeric(out_df["数量"], errors="coerce").fillna(0)
out_df["日時"] = pd.to_datetime(out_df["日時"], errors="coerce")
out_df = pd.merge(out_df, price_map, on="コード", how="left")
out_df["出庫金額"] = (out_df["数量"] * out_df["総単価"])
out_valid_date = (out_df.dropna(subset=["日時"]).copy())
out_valid_date["年月"] = (out_valid_date["日時"].dt.strftime("%Y-%m"))

month_list = sorted(out_valid_date["年月"].dropna().unique())


# ======== コード別集計 ========
code_group_columns = ["コード", "備品名", "単位"]
code_total = (
    out_df.groupby(code_group_columns, as_index=False, dropna=False)
    .agg(総消耗数=("数量", "sum"))
)

code_total = pd.merge(code_total, price_map, on="コード", how="left")

code_total["総金額"] = (code_total["総消耗数"] * code_total["総単価"])

code_monthly_quantity = pd.pivot_table(
    out_valid_date,
    index=code_group_columns,
    columns="年月",
    values="数量",
    aggfunc="sum",
    fill_value=0
)

code_monthly_amount = pd.pivot_table(
    out_valid_date,
    index=code_group_columns,
    columns="年月",
    values="出庫金額",
    aggfunc="sum",
    fill_value=0
)


if len(code_monthly_quantity.index) > 0:
    code_monthly_wide = pd.DataFrame(index=code_monthly_quantity.index)
else:
    code_monthly_wide = pd.DataFrame(columns=code_group_columns)

for month in month_list:
    if month in code_monthly_quantity.columns:
        code_monthly_wide[f"{month}_数量"] = code_monthly_quantity[month]
    else:
        code_monthly_wide[f"{month}_数量"] = np.nan

    if month in code_monthly_amount.columns:
        code_monthly_wide[f"{month}_金額"] = code_monthly_amount[month]
    else:
        code_monthly_wide[f"{month}_金額"] = np.nan

if len(code_monthly_quantity.index) > 0:
    code_monthly_wide = (code_monthly_wide.reset_index())


if month_list:
    code_result = pd.merge(code_total, code_monthly_wide, on=code_group_columns, how="outer")
else:
    code_result = code_total.copy()

code_fixed_columns = ["コード", "備品名", "総消耗数", "単位", "総単価", "総金額"]
code_monthly_columns = []
for month in month_list:
    code_monthly_columns.extend([f"{month}_数量", f"{month}_金額"])
code_result = code_result[code_fixed_columns + code_monthly_columns]
code_result = code_result.sort_values("コード", na_position="last")

code_total_row = {
    "コード": "合計",
    "備品名": "",
    "総消耗数": "",
    "単位": "",
    "総単価": "",
    "総金額": pd.to_numeric(code_result["総金額"], errors="coerce").fillna(0).sum()
}

for month in month_list:
    quantity_column = f"{month}_数量"
    amount_column = f"{month}_金額"
    code_total_row[quantity_column] = ""
    code_total_row[amount_column] = pd.to_numeric(code_result[amount_column], errors="coerce").fillna(0).sum()

code_result = pd.concat([code_result,pd.DataFrame([code_total_row])],ignore_index=True)

for month in month_list:
    quantity_column = f"{month}_数量"
    amount_column = f"{month}_金額"
    code_result[quantity_column] = (code_result[quantity_column].replace(0, np.nan))
    code_result[amount_column] = (code_result[amount_column].replace(0, np.nan))

department_group_columns = ["納品先", "コード", "備品名", "単位"]

department_total = (
    out_df.groupby(department_group_columns, as_index=False, dropna=False)
    .agg(総消耗数=("数量", "sum"))
)

department_total = pd.merge(department_total, price_map, on="コード", how="left")

department_total["総金額"] = (department_total["総消耗数"] * department_total["総単価"])


# ======== 部門別集計 ========
department_monthly_quantity = pd.pivot_table(
    out_valid_date,
    index=department_group_columns,
    columns="年月",
    values="数量",
    aggfunc="sum",
    fill_value=0
)

department_monthly_amount = pd.pivot_table(
    out_valid_date,
    index=department_group_columns,
    columns="年月",
    values="出庫金額",
    aggfunc="sum",
    fill_value=0
)

if len(department_monthly_quantity.index) > 0:
    department_monthly_wide = pd.DataFrame(index=department_monthly_quantity.index)
else:
    department_monthly_wide = pd.DataFrame(columns=department_group_columns)

for month in month_list:
    if month in department_monthly_quantity.columns:
        department_monthly_wide[f"{month}_数量"] = department_monthly_quantity[month]
    else:
        department_monthly_wide[f"{month}_数量"] = np.nan

    if month in department_monthly_amount.columns:
        department_monthly_wide[f"{month}_金額"] = department_monthly_amount[month]
    else:
        department_monthly_wide[f"{month}_金額"] = np.nan

if len(department_monthly_quantity.index) > 0:
    department_monthly_wide = (department_monthly_wide.reset_index())

if month_list:
    department_result = pd.merge(
        department_total,
        department_monthly_wide,
        on=department_group_columns,
        how="outer"
    )
else:
    department_result = department_total.copy()

department_fixed_columns = ["納品先", "コード", "備品名", "総消耗数", "単位", "総単価", "総金額"]

department_monthly_columns = []

for month in month_list:
    department_monthly_columns.extend([f"{month}_数量", f"{month}_金額"])

department_result = department_result[department_fixed_columns + department_monthly_columns]
department_result = (department_result.sort_values(["納品先", "コード"], na_position="last").reset_index(drop=True))

department_result_parts = []
for (department_name, department_group) in department_result.groupby("納品先", sort=False, dropna=False):
    department_group = (department_group.copy())
    department_result_parts.append(department_group)

    subtotal_row = {
        "納品先": department_name,
        "コード": "小計",
        "備品名": "",
        "総消耗数": "",
        "単位": "",
        "総単価": "",
        "総金額": pd.to_numeric(department_group["総金額"], errors="coerce").fillna(0).sum()
    }

    for month in month_list:
        quantity_column = (f"{month}_数量")
        amount_column = (f"{month}_金額")
        subtotal_row[quantity_column] = ""
        subtotal_row[amount_column] = pd.to_numeric(department_group[amount_column], errors="coerce").fillna(0).sum()

    department_result_parts.append(pd.DataFrame([subtotal_row]))

if department_result_parts:
    department_result = pd.concat(
        department_result_parts,
        ignore_index=True
    )

for month in month_list:
    quantity_column = f"{month}_数量"
    amount_column = f"{month}_金額"
    department_result[quantity_column] = (department_result[quantity_column].replace(0, np.nan))
    department_result[amount_column] = (department_result[amount_column].replace(0, np.nan))


# ======== Excel出力 ========
with pd.ExcelWriter(out_file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
    code_result.to_excel(writer, sheet_name=code_sheet_name, index=False)
    department_result.to_excel(writer, sheet_name=department_sheet_name, index=False)


# ======== Excelスタイル：コード別 ========
wb = load_workbook(out_file_path)
ws = wb[code_sheet_name]
ws.insert_rows(2)

code_headers = ["コード", "備品名", "総消耗数", "単位", "総単価", "総金額"]

for column_number, header in enumerate(code_headers, start=1):
    ws.cell(row=1, column=column_number, value=header)
    ws.merge_cells(start_row=1, start_column=column_number, end_row=2, end_column=column_number)

current_column = 7
for month in month_list:
    ws.merge_cells(start_row=1, start_column=current_column, end_row=1, end_column=current_column + 1)
    ws.cell(row=1, column=current_column, value=month)
    ws.cell(row=2, column=current_column, value="数量")
    ws.cell(row=2, column=current_column + 1, value="金額")
    current_column += 2

for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for row_number in range(3, ws.max_row + 1):
    for column_number in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_number, column=column_number)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

    ws.cell(row=row_number, column=1).number_format = "@"
    ws.cell(row=row_number, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=3).number_format = "#,##0.##"
    ws.cell(row=row_number, column=3).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=4).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=5).number_format = "#,##0.0"
    ws.cell(row=row_number, column=5).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=6).number_format = "#,##0"
    ws.cell(row=row_number, column=6).alignment = Alignment(horizontal="right", vertical="center")

    current_column = 7
    for month in month_list:
        ws.cell(row=row_number, column=current_column).number_format = "#,##0.##"
        ws.cell(row=row_number, column=current_column).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row_number, column=current_column + 1).number_format = "#,##0"
        ws.cell(row=row_number, column=current_column + 1).alignment = Alignment(horizontal="right", vertical="center")
        current_column += 2

code_widths = {1: 12, 2: 25, 3: 9, 4: 6, 5: 9, 6: 11}

for column_number, width in (code_widths.items()):
    ws.column_dimensions[get_column_letter(column_number)].width = width

current_column = 7

for month in month_list:
    ws.column_dimensions[get_column_letter(current_column)].width = 7
    ws.column_dimensions[get_column_letter(current_column + 1)].width = 9
    current_column += 2

code_last_row = ws.max_row

for cell in ws[code_last_row]:
    cell.font = total_font
    cell.fill = total_fill
    cell.border = thin_border
    cell.alignment = Alignment(vertical="center")

ws.cell(row=code_last_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
ws.cell(row=code_last_row, column=6).number_format = "#,##0"
ws.cell(row=code_last_row, column=6).alignment = Alignment(horizontal="right", vertical="center")

current_column = 7

for month in month_list:
    ws.cell(row=code_last_row, column=current_column + 1).number_format = "#,##0"
    ws.cell(row=code_last_row, column=current_column + 1).alignment = Alignment(horizontal="right", vertical="center")
    current_column += 2

ws.row_dimensions[1].height = 24
ws.row_dimensions[2].height = 22

for row_number in range(3, ws.max_row + 1):
    ws.row_dimensions[row_number].height = 20

ws.row_dimensions[code_last_row].height = 22
ws.freeze_panes = "G3"
ws.auto_filter.ref = None
ws.sheet_view.showGridLines = False


# ======== Excelスタイル：部門別 ========
ws = wb[department_sheet_name]
ws.insert_rows(2)
department_headers = ["納品先", "コード", "備品名", "総消耗数", "単位", "総単価", "総金額"]

for column_number, header in enumerate(department_headers, start=1):
    ws.cell(row=1, column=column_number, value=header)
    ws.merge_cells(start_row=1, start_column=column_number, end_row=2, end_column=column_number)

current_column = 8
for month in month_list:
    ws.merge_cells(start_row=1, start_column=current_column, end_row=1, end_column=current_column + 1)
    ws.cell(row=1, column=current_column, value=month)
    ws.cell(row=2, column=current_column, value="数量")
    ws.cell(row=2, column=current_column + 1, value="金額")
    current_column += 2

for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for row_number in range(3, ws.max_row + 1):
    for column_number in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_number, column=column_number)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

    ws.cell(row=row_number, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=2).number_format = "@"
    ws.cell(row=row_number, column=2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=3).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=4).number_format = "#,##0.##"
    ws.cell(row=row_number, column=4).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=5).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=6).number_format = "#,##0.0"
    ws.cell(row=row_number, column=6).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=7).number_format = "#,##0"
    ws.cell(row=row_number, column=7).alignment = Alignment(horizontal="right", vertical="center")

    current_column = 8
    for month in month_list:
        ws.cell(row=row_number, column=current_column).number_format = "#,##0.##"
        ws.cell(row=row_number, column=current_column).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row_number, column=current_column + 1).number_format = "#,##0"
        ws.cell(row=row_number, column=current_column + 1).alignment = Alignment(horizontal="right", vertical="center")
        current_column += 2

data_start_row = 3
row_number = data_start_row

while row_number <= ws.max_row:
    code_value = ws.cell(row=row_number, column=2).value
    if code_value == "小計":
        row_number += 1
        continue

    current_department = ws.cell(row=row_number, column=1).value

    merge_start_row = row_number
    merge_end_row = row_number

    check_row = row_number + 1

    while check_row <= ws.max_row:
        next_code = ws.cell(row=check_row, column=2).value
        next_department = ws.cell(row=check_row, column=1).value
        if next_code == "小計":
            break
        if (next_department != current_department):
            break

        merge_end_row = check_row
        check_row += 1

    if merge_end_row > merge_start_row:
        ws.merge_cells(start_row=merge_start_row, start_column=1, end_row=merge_end_row, end_column=1)
        ws.cell(row=merge_start_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

    row_number = check_row

for row_number in range(3, ws.max_row + 1):
    code_value = ws.cell(row=row_number, column=2).value

    if code_value == "小計":
        for column_number in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_number, column=column_number)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

        ws.cell(row=row_number, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row_number, column=2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row_number, column=7).number_format = "#,##0"
        ws.cell(row=row_number, column=7).alignment = Alignment(horizontal="right", vertical="center")

        current_column = 8
        for month in month_list:
            ws.cell(row=row_number, column=current_column).alignment = Alignment(horizontal="right", vertical="center")
            ws.cell(row=row_number, column=current_column + 1).number_format = "#,##0"
            ws.cell(row=row_number, column=current_column + 1).alignment = Alignment(horizontal="right", vertical="center")
            current_column += 2

        ws.row_dimensions[row_number].height = 22

department_widths = {1: 10, 2: 12, 3: 25, 4: 9, 5: 6, 6: 9, 7: 11}

for column_number, width in (department_widths.items()):
    ws.column_dimensions[get_column_letter(column_number)].width = width

current_column = 8
for month in month_list:
    ws.column_dimensions[get_column_letter(current_column)].width = 7
    ws.column_dimensions[get_column_letter(current_column + 1)].width = 9
    current_column += 2

ws.row_dimensions[1].height = 24
ws.row_dimensions[2].height = 22

for row_number in range(3, ws.max_row + 1):
    if (ws.cell(row=row_number, column=2).value != "小計"):
        ws.row_dimensions[row_number].height = 20

ws.freeze_panes = "H3"

ws.auto_filter.ref = None
ws.sheet_view.showGridLines = False

wb.save(out_file_path)

print("完了")