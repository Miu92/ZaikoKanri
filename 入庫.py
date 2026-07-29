import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side
)
from openpyxl.utils import get_column_letter


file_path = r"\\dionas\管理部\備品管理\入出庫履歴\2026入庫履歴.xlsx"
sheet_name = "コード別集計"

df = pd.read_excel(file_path)
df["金額"] = pd.to_numeric(df["金額"], errors="coerce").fillna(0)
df["数量"] = pd.to_numeric(df["数量"], errors="coerce").fillna(0)
df["コード"] = df["コード"].astype(str)
df["日時"] = pd.to_datetime(df["日時"], errors="coerce")

group_columns = ["コード", "備品名", "単位"]


# ======== 年間：数量・単価・金額計算 ========
total_summary = (
    df.groupby(group_columns, as_index=False, dropna=False)
    .agg({"数量": "sum", "金額": "sum"})
)
total_summary["総単価"] = np.where(
    total_summary["数量"] != 0,
    total_summary["金額"] / total_summary["数量"],
    np.nan
)
total_summary["総単価"] = total_summary["総単価"].round(1)
total_summary = total_summary.rename(columns={"数量": "総数量", "金額": "総金額"})
total_summary = total_summary[["コード", "備品名", "単位", "総数量", "総単価", "総金額"]]


# ======== 月別：数量・単価・金額計算 ========
df_monthly = (df.dropna(subset=["日時"]).copy())
df_monthly["年月"] = (df_monthly["日時"].dt.strftime("%Y-%m"))

monthly_summary = (
    df_monthly.groupby(group_columns + ["年月"], as_index=False, dropna=False)
    .agg({"数量": "sum", "金額": "sum"})
)
monthly_summary["単価"] = np.where(
    monthly_summary["数量"] != 0,
    monthly_summary["金額"] / monthly_summary["数量"],
    np.nan
)
monthly_summary["単価"] = (monthly_summary["単価"].round(1))


# ======== 月別：横並び ========
month_list = sorted(monthly_summary["年月"].dropna().unique())

quantity_pivot = monthly_summary.pivot_table(
    index=group_columns,
    columns="年月",
    values="数量",
    aggfunc="sum"
)

unit_price_pivot = monthly_summary.pivot_table(
    index=group_columns,
    columns="年月",
    values="単価",
    aggfunc="first"
)

amount_pivot = monthly_summary.pivot_table(
    index=group_columns,
    columns="年月",
    values="金額",
    aggfunc="sum"
)

monthly_wide = pd.DataFrame(index=quantity_pivot.index)

for month in month_list:
    monthly_wide[f"{month}_数量"] = (
        quantity_pivot[month]
        if month in quantity_pivot.columns
        else np.nan
    )

    monthly_wide[f"{month}_単価"] = (
        unit_price_pivot[month]
        if month in unit_price_pivot.columns
        else np.nan
    )

    monthly_wide[f"{month}_金額"] = (
        amount_pivot[month]
        if month in amount_pivot.columns
        else np.nan
    )

monthly_wide = monthly_wide.reset_index()


# ======== 合併 ========
result = pd.merge(total_summary, monthly_wide, on=group_columns, how="outer")
fixed_columns = ["コード", "備品名", "総数量", "単位", "総単価", "総金額"]
monthly_columns = []

for month in month_list:
    monthly_columns.extend([f"{month}_数量", f"{month}_単価", f"{month}_金額"])

result = result[fixed_columns + monthly_columns]
result = result.sort_values("コード", na_position="last")


# ======== 合計金額 ========
total_row = {
    "コード": "合計",
    "備品名": "",
    "総数量": "",
    "単位": "",
    "総単価": "",
    "総金額": pd.to_numeric(
        result["総金額"],
        errors="coerce"
    ).fillna(0).sum()
}

for month in month_list:
    quantity_column = f"{month}_数量"
    price_column = f"{month}_単価"
    amount_column = f"{month}_金額"

    total_row[quantity_column] = ""
    total_row[price_column] = ""

    total_row[amount_column] = pd.to_numeric(result[amount_column], errors="coerce").fillna(0).sum()

result = pd.concat([result, pd.DataFrame([total_row])], ignore_index=True)


# ======== 0⇒　 ========
for month in month_list:
    quantity_column = f"{month}_数量"
    price_column = f"{month}_単価"
    amount_column = f"{month}_金額"

    result[quantity_column] = result[quantity_column].replace(0, np.nan)
    result[price_column] = result[price_column].replace(0, np.nan)
    result[amount_column] = result[amount_column].replace(0, np.nan)


# ======== Excel出力 ========
with (pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer):
    result.to_excel(writer, sheet_name=sheet_name, index=False)


# ======== Excelスタイル ========
wb = load_workbook(file_path)
ws = wb[sheet_name]
ws.insert_rows(2)

fixed_headers = ["コード", "備品名", "総数量", "単位", "総単価", "総金額"]
for column_number, header in enumerate(fixed_headers, start=1):
    ws.cell(row=1, column=column_number, value=header)
    ws.merge_cells(start_row=1, start_column=column_number, end_row=2, end_column=column_number)

start_column = len(fixed_headers) + 1
for month in month_list:
    month_start_column = start_column
    month_end_column = start_column + 2
    ws.merge_cells(start_row=1, start_column=month_start_column, end_row=1, end_column=month_end_column)
    ws.cell(row=1, column=month_start_column, value=month)
    ws.cell(row=2, column=month_start_column, value="数量")
    ws.cell(row=2, column=month_start_column + 1, value="単価")
    ws.cell(row=2, column=month_start_column + 2, value="金額")
    start_column += 3

header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
header_font = Font(bold=True)
total_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
total_font = Font(bold=True)
thin_side = Side(style="thin", color="808080")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1,max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

for row_number in range(3, ws.max_row + 1):
    ws.cell(row=row_number, column=1).number_format = "@"
    ws.cell(row=row_number, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=3).number_format = "#,##0.##"
    ws.cell(row=row_number, column=3).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=4).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row_number, column=5).number_format = "#,##0.0"
    ws.cell(row=row_number, column=5).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=row_number, column=6).number_format = "#,##0"
    ws.cell(row=row_number, column=6).alignment = Alignment(horizontal="right", vertical="center")

    current_column = 7
    for month in month_list:
        ws.cell(row=row_number, column=current_column).number_format = "#,##0.##"
        ws.cell(row=row_number, column=current_column).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row_number, column=current_column + 1).number_format = "#,##0.0"
        ws.cell(row=row_number, column=current_column + 1).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row_number, column=current_column + 2).number_format = "#,##0"
        ws.cell(row=row_number, column=current_column + 2).alignment = Alignment(horizontal="right", vertical="center")
        current_column += 3

column_widths = {1: 12, 2: 25, 3: 9, 4: 6, 5: 9, 6: 11}

for column_number, width in column_widths.items():
    ws.column_dimensions[get_column_letter(column_number)].width = width

current_column = 7
for month in month_list:
    ws.column_dimensions[get_column_letter(current_column)].width = 7
    ws.column_dimensions[get_column_letter(current_column + 1)].width = 8
    ws.column_dimensions[get_column_letter(current_column + 2)].width = 9
    current_column += 3

last_row = ws.max_row
for cell in ws[last_row]:
    cell.font = total_font
    cell.fill = total_fill
    cell.border = thin_border
    cell.alignment = Alignment(vertical="center")

ws.cell(row=last_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
ws.cell(row=last_row, column=2).alignment = Alignment(horizontal="left", vertical="center")
ws.cell(row=last_row, column=4).alignment = Alignment(horizontal="left", vertical="center")
ws.cell(row=last_row, column=6).number_format = "#,##0"
ws.cell(row=last_row, column=6).alignment = Alignment(horizontal="right", vertical="center")

current_column = 7
for month in month_list:
    amount_column = current_column + 2
    ws.cell(row=last_row, column=amount_column).number_format = "#,##0"
    ws.cell(row=last_row, column=amount_column).alignment = Alignment(horizontal="right", vertical="center")
    current_column += 3

ws.row_dimensions[last_row].height = 22

ws.freeze_panes = "G3"

ws.row_dimensions[1].height = 24
ws.row_dimensions[2].height = 22

for row_number in range(3,ws.max_row):
    ws.row_dimensions[row_number].height = 20

ws.auto_filter.ref = None
ws.sheet_view.showGridLines = False

wb.save(file_path)

print("完了")
